import hashlib
import hmac
import json
import logging
import os
import re
import sqlite3
import tempfile
import time
from contextlib import closing
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import quote
from zipfile import BadZipFile

import requests
from dotenv import load_dotenv
from fastapi import APIRouter, Depends, FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel

load_dotenv()

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"
DATABASE_PATH = Path(os.getenv("DATABASE_PATH", str(BASE_DIR / "data" / "products.sqlite3")))
EPHEMERAL_DATABASE_PATH = Path(tempfile.gettempdir()) / "oslo_salla_products.sqlite3"
SKU_PREFIX = "PRD"
DEFAULT_QTY = 100
DEFAULT_MAX_QTY = 10
DEFAULT_BRAND = "برايز بوكس | Prize Box"
PRICE_MARKUP = 100
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")
GEMINI_URL = (
    "https://generativelanguage.googleapis.com/v1beta/models/"
    f"{GEMINI_MODEL}:generateContent"
)

SALLA_COLUMNS = [
    "النوع ",
    "أسم المنتج",
    "تصنيف المنتج",
    "صورة المنتج",
    "وصف صورة المنتج",
    "نوع المنتج",
    "سعر المنتج",
    "الوصف",
    "هل يتطلب شحن؟",
    "رمز المنتج sku",
    "سعر التكلفة",
    "السعر المخفض",
    "تاريخ بداية التخفيض",
    "تاريخ نهاية التخفيض",
    "اقصي كمية لكل عميل",
    "إخفاء خيار تحديد الكمية",
    "اضافة صورة عند الطلب",
    "الوزن",
    "وحدة الوزن",
    "الماركة",
    "العنوان الترويجي",
    "تثبيت المنتج",
    "الباركود",
    "السعرات الحرارية",
    "MPN",
    "GTIN",
    "خاضع للضريبة ؟",
    "سبب عدم الخضوع للضريبة",
    "[1] الاسم",
    "[1] النوع",
    "[1] القيمة",
    "[1] الصورة / اللون",
    "[2] الاسم",
    "[2] النوع",
    "[2] القيمة",
    "[2] الصورة / اللون",
    "[3] الاسم",
    "[3] النوع",
    "[3] القيمة",
    "[3] الصورة / اللون",
]
EXTRA_COLUMNS = ["الكمية المتوفرة", "عنوان SEO", "وصف SEO"]
ALL_COLUMNS = SALLA_COLUMNS + EXTRA_COLUMNS

GEMINI_SYSTEM_PROMPT = """أنت خبير تسويق إلكتروني ومتخصص SEO لمتاجر سلة (Salla) في السعودية والخليج.
ستُعطى وصفاً حراً أو ملاحظات غير منظمة عن منتج من تاجر. مهمتك استخراج/توليد بيانات
المنتج وإعادتها بصيغة JSON فقط (بدون أي نص أو شرح أو Markdown قبل أو بعد الـ JSON).

التزم بالتالي:
- اسم المنتج: واضح، جذاب، يحتوي على أهم كلمة مفتاحية يبحث بها العميل، بدون حشو كلمات.
- الوصف: فقرة تسويقية غنية بالكلمات المفتاحية ذات الصلة الطبيعية (بدون حشو مبالغ فيه)
  تشرح فوائد المنتج ومن يناسبه، وتُحفّز على الشراء، ولا تكون مجرد نسخ لكلام المستخدم.
- عنوان SEO: عنوان قصير (يُفضّل بين 50 و60 حرفاً) يحتوي على الكلمة المفتاحية الأهم
  وغالباً اسم العلامة/المتجر إن وُجد ضمن النص، مناسب لعرضه كعنوان نتيجة بحث جوجل.
- وصف SEO: وصف Meta قصير (يُفضّل بين 120 و160 حرفاً) يلخّص قيمة المنتج ويحتوي
  كلمات مفتاحية طبيعية ويشجّع المستخدم على النقر من نتائج البحث.
- التصنيف: تصنيف مقترح للمنتج بصيغة "قسم > قسم فرعي" إن أمكن استنتاجه، وإلا اتركه فارغاً.
- الماركة: استخرجها فقط إن ذُكرت صريحاً في النص، وإلا اتركها فارغة.
- وصف الصورة (alt text): جملة قصيرة وصفية للصورة تخدم السيو، تتضمن اسم المنتج.
- العنوان الترويجي: جملة تسويقية قصيرة جداً (Badge) مثل "الأكثر مبيعاً" أو "عرض محدود"
  فقط إذا كان مناسباً من سياق النص، وإلا اتركه فارغاً.
- السعر وسعر التكلفة والسعر المخفض: استخرجها كأرقام فقط إن وُجدت أرقام صريحة في النص
  (بدون رمز عملة)، وإلا اتركها فارغة (سلسلة نصية فارغة).
- لا تخترع معلومات غير موجودة وغير قابلة للاستنتاج المنطقي (مثل أرقام GTIN/MPN/الباركود)
  واتركها فارغة دوماً ما لم تُذكر صريحاً.
- اكتب كل المحتوى باللغة العربية الفصحى التسويقية السليمة ما لم يطلب المستخدم لغة أخرى.

أعد فقط الكائن JSON التالي بالضبط بهذه المفاتيح (بدون أي مفاتيح إضافية):
{
  "اسم_المنتج": "",
  "الوصف": "",
  "تصنيف_المنتج": "",
  "الماركة": "",
  "وصف_صورة_المنتج": "",
  "العنوان_الترويجي": "",
  "سعر_المنتج": "",
  "سعر_التكلفة": "",
  "السعر_المخفض": "",
  "عنوان_SEO": "",
  "وصف_SEO": ""
}"""

app = FastAPI(title="OSLO Salla Products")
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


def _db() -> sqlite3.Connection:
    global DATABASE_PATH
    try:
        DATABASE_PATH.parent.mkdir(parents=True, exist_ok=True)
        connection = sqlite3.connect(DATABASE_PATH, timeout=30)
    except PermissionError:
        if DATABASE_PATH == EPHEMERAL_DATABASE_PATH:
            raise
        logging.warning(
            "Database path %s is not writable; using temporary storage at %s. "
            "Data may be lost when the service restarts.",
            DATABASE_PATH,
            EPHEMERAL_DATABASE_PATH,
        )
        DATABASE_PATH = EPHEMERAL_DATABASE_PATH
        DATABASE_PATH.parent.mkdir(parents=True, exist_ok=True)
        connection = sqlite3.connect(DATABASE_PATH, timeout=30)
    except sqlite3.OperationalError as exc:
        if "unable to open database file" not in str(exc).lower():
            raise
        if DATABASE_PATH == EPHEMERAL_DATABASE_PATH:
            raise
        logging.warning(
            "Database path %s could not be opened; using temporary storage at %s. "
            "Data may be lost when the service restarts.",
            DATABASE_PATH,
            EPHEMERAL_DATABASE_PATH,
        )
        DATABASE_PATH = EPHEMERAL_DATABASE_PATH
        DATABASE_PATH.parent.mkdir(parents=True, exist_ok=True)
        connection = sqlite3.connect(DATABASE_PATH, timeout=30)
    connection.row_factory = sqlite3.Row
    return connection


def _init_db() -> None:
    with closing(_db()) as connection:
        connection.execute(
            """CREATE TABLE IF NOT EXISTS products (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                data TEXT NOT NULL,
                created_at TEXT NOT NULL
            )"""
        )
        connection.commit()


@app.on_event("startup")
def startup() -> None:
    _init_db()


def _session_secret() -> str:
    value = os.getenv("SESSION_SECRET", "")
    if len(value) < 32:
        raise HTTPException(
            status_code=503,
            detail="على مسؤول الاستضافة إعداد SESSION_SECRET بطول 32 محرفًا على الأقل.",
        )
    return value


def _cookie_secure() -> bool:
    return os.getenv("COOKIE_SECURE", "true").strip().lower() not in {"0", "false", "no"}


def _valid_session(request: Request) -> bool:
    token = request.cookies.get("oslo_session", "")
    try:
        expires_text, supplied_signature = token.split(".", 1)
        expires = int(expires_text)
    except (ValueError, TypeError):
        return False
    if expires < int(time.time()):
        return False
    expected = hmac.new(
        _session_secret().encode(), expires_text.encode(), hashlib.sha256
    ).hexdigest()
    return hmac.compare_digest(supplied_signature, expected)


def require_auth(request: Request) -> None:
    if not _valid_session(request):
        raise HTTPException(status_code=401, detail="يرجى تسجيل الدخول أولاً.")


protected = APIRouter(prefix="/api", dependencies=[Depends(require_auth)])


class LoginBody(BaseModel):
    password: str


class AIRequest(BaseModel):
    text: str


class ProductRequest(BaseModel):
    data: dict[str, Any]


class DeleteRequest(BaseModel):
    ids: list[int]


class SkuRequest(BaseModel):
    sku: str


@app.get("/api/me")
def me(request: Request) -> dict[str, bool]:
    return {"authenticated": _valid_session(request)}


@app.post("/api/login")
def login(body: LoginBody, request: Request):
    configured_password = os.getenv("APP_PASSWORD", "")
    if not configured_password:
        raise HTTPException(
            status_code=503,
            detail="على مسؤول الاستضافة إعداد APP_PASSWORD قبل استخدام الموقع.",
        )
    _session_secret()
    if not hmac.compare_digest(body.password, configured_password):
        raise HTTPException(status_code=401, detail="كلمة المرور غير صحيحة.")
    expires = str(int(time.time()) + 60 * 60 * 24 * 7)
    signature = hmac.new(
        _session_secret().encode(), expires.encode(), hashlib.sha256
    ).hexdigest()
    response = {"ok": True}
    from fastapi.responses import JSONResponse

    result = JSONResponse(response)
    result.set_cookie(
        "oslo_session",
        f"{expires}.{signature}",
        httponly=True,
        secure=_cookie_secure(),
        samesite="strict",
        max_age=60 * 60 * 24 * 7,
        path="/",
    )
    return result


@app.post("/api/logout")
def logout():
    from fastapi.responses import JSONResponse

    response = JSONResponse({"ok": True})
    response.delete_cookie("oslo_session", path="/", samesite="strict")
    return response


def _next_sku(connection: sqlite3.Connection) -> str:
    highest = 0
    for record in connection.execute("SELECT data FROM products"):
        data = json.loads(record["data"])
        match = re.search(r"(\d+)$", str(data.get("رمز المنتج sku", "")))
        if match:
            highest = max(highest, int(match.group(1)))
    return f"{SKU_PREFIX}-{highest + 1:03d}"


def _product_count(connection: sqlite3.Connection) -> int:
    return int(connection.execute("SELECT COUNT(*) FROM products").fetchone()[0])


def _is_ephemeral_storage() -> bool:
    return DATABASE_PATH.resolve().is_relative_to(Path(tempfile.gettempdir()).resolve())


@protected.get("/config")
def config_status():
    cloudinary_ready = all(
        os.getenv(name)
        for name in (
            "CLOUDINARY_CLOUD_NAME",
            "CLOUDINARY_API_KEY",
            "CLOUDINARY_API_SECRET",
        )
    )
    return {
        "gemini_ready": bool(os.getenv("GEMINI_API_KEY")),
        "cloudinary_ready": cloudinary_ready,
        "storage_ephemeral": _is_ephemeral_storage(),
    }


@protected.get("/products")
def list_products():
    with closing(_db()) as connection:
        records = connection.execute(
            "SELECT id, data FROM products ORDER BY id"
        ).fetchall()
        products = [{"id": row["id"], "data": json.loads(row["data"])} for row in records]
        return {
            "products": products,
            "count": len(products),
            "next_sku": _next_sku(connection),
        }


@protected.post("/products")
def create_product(body: ProductRequest):
    data = {column: body.data.get(column, "") for column in ALL_COLUMNS}
    name = str(data["أسم المنتج"]).strip()
    if not name:
        raise HTTPException(status_code=422, detail="يرجى إدخال اسم المنتج.")
    try:
        data["سعر المنتج"] = float(data["سعر المنتج"])
    except (TypeError, ValueError):
        raise HTTPException(status_code=422, detail="سعر المنتج يجب أن يكون رقمًا.")
    if not data["رمز المنتج sku"]:
        data["رمز المنتج sku"] = ""
    data["النوع "] = "منتج"
    data["الكمية المتوفرة"] = (
        data["الكمية المتوفرة"] if data["الكمية المتوفرة"] != "" else DEFAULT_QTY
    )
    with closing(_db()) as connection:
        try:
            connection.execute("BEGIN IMMEDIATE")
            if not data["رمز المنتج sku"]:
                data["رمز المنتج sku"] = _next_sku(connection)
            connection.execute(
                "INSERT INTO products (data, created_at) VALUES (?, ?)",
                (json.dumps(data, ensure_ascii=False), datetime.now().isoformat()),
            )
            total = _product_count(connection)
            next_sku = _next_sku(connection)
            connection.commit()
        except Exception:
            connection.rollback()
            raise
    return {"ok": True, "count": total, "next_sku": next_sku, "sku": data["رمز المنتج sku"]}


@protected.put("/products/{product_id}/sku")
def update_sku(product_id: int, body: SkuRequest):
    sku = body.sku.strip()
    if not sku:
        raise HTTPException(status_code=422, detail="أدخل رمز SKU.")
    with closing(_db()) as connection:
        connection.execute("BEGIN IMMEDIATE")
        record = connection.execute(
            "SELECT data FROM products WHERE id = ?", (product_id,)
        ).fetchone()
        if record is None:
            connection.rollback()
            raise HTTPException(status_code=404, detail="المنتج غير موجود.")
        data = json.loads(record["data"])
        data["رمز المنتج sku"] = sku
        connection.execute(
            "UPDATE products SET data = ? WHERE id = ?",
            (json.dumps(data, ensure_ascii=False), product_id),
        )
        connection.commit()
    return {"ok": True}


@protected.delete("/products")
def delete_products(body: DeleteRequest):
    if not body.ids:
        raise HTTPException(status_code=422, detail="حدد منتجًا واحدًا على الأقل.")
    ids = list(dict.fromkeys(body.ids))
    with closing(_db()) as connection:
        connection.executemany("DELETE FROM products WHERE id = ?", [(item,) for item in ids])
        count = _product_count(connection)
        next_sku = _next_sku(connection)
        connection.commit()
    return {"ok": True, "deleted": len(ids), "count": count, "next_sku": next_sku}


@protected.post("/products/renumber")
def renumber_products():
    with closing(_db()) as connection:
        connection.execute("BEGIN IMMEDIATE")
        records = connection.execute("SELECT id, data FROM products ORDER BY id").fetchall()
        for index, record in enumerate(records, start=1):
            data = json.loads(record["data"])
            data["رمز المنتج sku"] = f"{SKU_PREFIX}-{index:03d}"
            connection.execute(
                "UPDATE products SET data = ? WHERE id = ?",
                (json.dumps(data, ensure_ascii=False), record["id"]),
            )
        connection.commit()
        next_sku = _next_sku(connection)
    return {"ok": True, "count": len(records), "next_sku": next_sku}


def _build_workbook(records: list[dict[str, Any]]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Salla Products Template Sheet"
    header_fill = PatternFill("solid", fgColor="5DD5C4")
    extra_fill = PatternFill("solid", fgColor="FBBF24")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
    n_salla = len(SALLA_COLUMNS)
    n_total = len(ALL_COLUMNS)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_salla)
    cell = sheet.cell(1, 1, "بيانات المنتج (متوافق مع قالب سلة الرسمي)")
    cell.fill = header_fill
    cell.font = Font(bold=True, size=14, name="Arial")
    cell.alignment = center
    sheet.row_dimensions[1].height = 30
    sheet.merge_cells(start_row=1, start_column=n_salla + 1, end_row=1, end_column=n_total)
    cell = sheet.cell(
        1, n_salla + 1, "⚠ أعمدة إضافية (للاستخدام الداخلي فقط - ليست من قالب سلة)"
    )
    cell.fill = extra_fill
    cell.font = Font(bold=True, size=12, name="Arial")
    cell.alignment = center
    for column_index, column_name in enumerate(ALL_COLUMNS, start=1):
        cell = sheet.cell(2, column_index, column_name)
        cell.fill = header_fill if column_index <= n_salla else extra_fill
        cell.font = Font(bold=True, size=14, name="Arial")
        cell.alignment = center
    sheet.row_dimensions[2].height = 35
    widths = {
        1: 12, 2: 25, 3: 22, 4: 45, 5: 22, 6: 16, 7: 14, 8: 30,
        9: 16, 10: 16, 11: 14, 12: 14, 13: 22, 14: 22, 15: 18,
        16: 22, 17: 20, 18: 10, 19: 12, 20: 16, 21: 18, 22: 14,
        23: 18, 24: 20, 25: 10, 26: 10, 27: 16, 28: 30, 29: 16,
        30: 12, 31: 30, 32: 22, 33: 16, 34: 12, 35: 30, 36: 22,
        37: 16, 38: 12, 39: 30, 40: 22, 41: 16, 42: 30, 43: 40,
    }
    for column_index, width in widths.items():
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    for row_index, data in enumerate(records, start=3):
        for column_index, name in enumerate(ALL_COLUMNS, start=1):
            cell = sheet.cell(row_index, column_index, data.get(name, ""))
            cell.alignment = Alignment(
                horizontal="right", vertical="center", wrap_text=True, readingOrder=2
            )
            cell.font = Font(size=11, name="Arial")
        sheet.row_dimensions[row_index].height = 22
    categories = workbook.create_sheet("Categories, Types & Brands")
    categories.cell(1, 1, "ما يمتلكه متجرك الحالى من تصنيفات")
    categories.cell(1, 2, "أنواع المنتجات المتاحة حالياَ")
    categories.cell(1, 3, "ما يمتلكه متجرك الحالى من ماركات تجارية")
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _workbook_response(records: list[dict[str, Any]]):
    output = _build_workbook(records)
    filename = quote("منتجات_سلة.xlsx")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


@protected.get("/export.xlsx")
def export_products():
    with closing(_db()) as connection:
        records = [
            json.loads(row["data"])
            for row in connection.execute("SELECT data FROM products ORDER BY id")
        ]
    return _workbook_response(records)


@protected.get("/template.xlsx")
def download_template():
    return _workbook_response([])


@protected.post("/import.xlsx")
def import_products(file: UploadFile = File(...)):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=422, detail="اختر ملف Excel بصيغة .xlsx.")
    try:
        workbook = load_workbook(file.file, read_only=True, data_only=True)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[2]]
        required = {"النوع ", "أسم المنتج", "سعر المنتج"}
        if not required.issubset(set(headers)):
            workbook.close()
            raise HTTPException(
                status_code=422,
                detail="لم يتم العثور على رؤوس أعمدة قالب سلة في الصف الثاني.",
            )
        columns = {header: index for index, header in enumerate(headers) if header}
        rows = []
        for values in sheet.iter_rows(min_row=3, values_only=True):
            if not any(value is not None and value != "" for value in values):
                continue
            data = {column: "" for column in ALL_COLUMNS}
            for column in ALL_COLUMNS:
                index = columns.get(column)
                if index is not None and index < len(values) and values[index] is not None:
                    value = values[index]
                    data[column] = value.isoformat() if isinstance(value, datetime) else value
            rows.append(data)
        workbook.close()
    except HTTPException:
        raise
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise HTTPException(status_code=422, detail="تعذرت قراءة ملف Excel المحدد.") from exc
    if not rows:
        raise HTTPException(status_code=422, detail="لا يحتوي ملف Excel على منتجات لاستيرادها.")
    with closing(_db()) as connection:
        connection.execute("BEGIN IMMEDIATE")
        for data in rows:
            if not data["رمز المنتج sku"]:
                data["رمز المنتج sku"] = _next_sku(connection)
            connection.execute(
                "INSERT INTO products (data, created_at) VALUES (?, ?)",
                (json.dumps(data, ensure_ascii=False, default=str), datetime.now().isoformat()),
            )
        total = _product_count(connection)
        next_sku = _next_sku(connection)
        connection.commit()
    return {"ok": True, "imported": len(rows), "count": total, "next_sku": next_sku}


@protected.post("/ai/extract")
def ai_extract(body: AIRequest):
    text = body.text.strip()
    if not text:
        raise HTTPException(status_code=422, detail="يرجى كتابة وصف المنتج أولاً.")
    api_key = os.getenv("GEMINI_API_KEY", "").strip()
    if not api_key:
        raise HTTPException(
            status_code=503,
            detail="ميزة الذكاء الاصطناعي غير مهيأة. أضف GEMINI_API_KEY إلى إعدادات الخادم.",
        )
    payload = {
        "contents": [{"role": "user", "parts": [{"text": text}]}],
        "systemInstruction": {"parts": [{"text": GEMINI_SYSTEM_PROMPT}]},
        "generationConfig": {"temperature": 0.4, "responseMimeType": "application/json"},
    }
    try:
        response = requests.post(
            GEMINI_URL,
            headers={"Content-Type": "application/json", "x-goog-api-key": api_key},
            json=payload,
            timeout=60,
        )
    except requests.RequestException as exc:
        raise HTTPException(status_code=502, detail=f"تعذر الاتصال بـ Gemini: {exc}") from exc
    if response.status_code != 200:
        raise HTTPException(
            status_code=502,
            detail=f"خطأ من Gemini ({response.status_code}): {response.text[:500]}",
        )
    try:
        result = response.json()
        text_result = result["candidates"][0]["content"]["parts"][0]["text"].strip()
        if text_result.startswith("```"):
            text_result = re.sub(r"^```[a-zA-Z]*\s*", "", text_result)
            text_result = re.sub(r"```\s*$", "", text_result)
        data = json.loads(text_result)
    except (ValueError, KeyError, IndexError, TypeError) as exc:
        raise HTTPException(
            status_code=502, detail="تعذر تحليل رد Gemini إلى بيانات JSON صالحة."
        ) from exc
    return {"data": data}


@protected.post("/images/upload")
def upload_image(file: UploadFile = File(...)):
    cloud_name = os.getenv("CLOUDINARY_CLOUD_NAME", "").strip()
    api_key = os.getenv("CLOUDINARY_API_KEY", "").strip()
    api_secret = os.getenv("CLOUDINARY_API_SECRET", "").strip()
    if not all((cloud_name, api_key, api_secret)):
        raise HTTPException(
            status_code=503,
            detail="رفع الصور غير مهيأ. أضف بيانات CLOUDINARY إلى إعدادات الخادم.",
        )
    timestamp = str(int(time.time()))
    signature = hashlib.sha1(
        f"timestamp={timestamp}{api_secret}".encode("utf-8")
    ).hexdigest()
    try:
        response = requests.post(
            f"https://api.cloudinary.com/v1_1/{cloud_name}/image/upload",
            files={
                "file": (
                    file.filename or "image",
                    file.file,
                    file.content_type or "application/octet-stream",
                )
            },
            data={"api_key": api_key, "timestamp": timestamp, "signature": signature},
            timeout=60,
        )
    except requests.RequestException as exc:
        raise HTTPException(status_code=502, detail=f"تعذر رفع الصورة: {exc}") from exc
    if response.status_code != 200:
        raise HTTPException(
            status_code=502,
            detail=f"خطأ Cloudinary ({response.status_code}): {response.text[:500]}",
        )
    try:
        url = response.json()["secure_url"]
    except (ValueError, KeyError, TypeError) as exc:
        raise HTTPException(status_code=502, detail="لم يُرجع Cloudinary رابط الصورة.") from exc
    return {"url": url, "filename": file.filename or ""}


app.include_router(protected)


@app.get("/")
def home():
    return FileResponse(STATIC_DIR / "index.html")
